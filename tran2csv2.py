import os
import glob
import openpyxl
import tkinter as tk
from tkinter import filedialog, messagebox, scrolledtext

class IVTransApp:
    def __init__(self, root):
        self.root = root
        self.root.title("HCI IV Data 轉檔工具")
        self.root.geometry("800x650")

        # 儲存轉換後的資料結構（格式：{檔名: 資料內容清單}）
        self.trans_results = {}

        self.create_widgets()

    def create_widgets(self):
        # 1. 路徑選取區塊
        path_frame = tk.LabelFrame(self.root, text=" 資料夾設定 ", font=("Arial", 12, "bold"))
        path_frame.pack(fill="x", padx=15, pady=10)

        self.path_entry = tk.Entry(path_frame, font=("Arial", 11))
        self.path_entry.pack(side="left", fill="x", expand=True, padx=10, pady=10)

        browse_btn = tk.Button(path_frame, text="瀏覽資料夾", font=("Arial", 10), command=self.browse_folder)
        browse_btn.pack(side="right", padx=10, pady=10)

        # 2. 功能按鈕區塊
        btn_frame = tk.Frame(self.root)
        btn_frame.pack(fill="x", padx=15, pady=5)

        self.btn_hci = tk.Button(btn_frame, text="HCI IV DATA", font=("Arial", 11, "bold"), 
                                 bg="#4CAF50", fg="white", width=20, height=2, command=self.run_hci_trans)
        self.btn_hci.pack(side="left", padx=10)

        self.btn_save = tk.Button(btn_frame, text="SAVE IV DATA", font=("Arial", 11, "bold"), 
                                  bg="#2196F3", fg="white", width=20, height=2, command=self.run_save_dat)
        self.btn_save.pack(side="right", padx=10)

        # 3. Log 記錄視窗區塊
        log_frame = tk.LabelFrame(self.root, text=" 操作記錄 (Log) ", font=("Arial", 12, "bold"))
        log_frame.pack(fill="both", expand=True, padx=15, pady=10)

        self.log_text = scrolledtext.ScrolledText(log_frame, font=("Consolas", 10), bg="#1e1e1e", fg="#00ff00")
        self.log_text.pack(fill="both", expand=True, padx=10, pady=10)

        self.log("系統就緒，請選擇要讀取的資料夾。\n")

    def log(self, message):
        """寫入 Log 訊息到介面"""
        self.log_text.insert(tk.END, message + "\n")
        self.log_text.see(tk.END)

    def browse_folder(self):
        """選擇資料夾按鈕事件"""
        folder_selected = filedialog.askdirectory()
        if folder_selected:
            self.path_entry.delete(0, tk.END)
            self.path_entry.insert(0, folder_selected)
            self.log(f"已選擇資料夾: {folder_selected}")
            
            # 列出資料夾內有哪些檔案
            files = os.listdir(folder_selected)
            self.log(f"--- 該資料夾內容清單 ({len(files)} 個項目) ---")
            for f in files:
                self.log(f"  • {f}")
            self.log("-" * 50)

    def get_cell_val(self, sheet, row, col):
        """直接從 openpyxl 取得儲存格的值，轉為字串並保持原貌（不進行任何數值轉型）"""
        val = sheet.cell(row=row, column=col).value
        if val is None:
            return ""
        return str(val).strip()

    def run_hci_trans(self):
        """執行 HCI IV DATA 轉檔邏輯"""
        target_dir = self.path_entry.get().strip()
        if not target_dir or not os.path.isdir(target_dir):
            messagebox.showerror("錯誤", "請先選擇有效的資料夾路徑！")
            return

        self.log("=== 開始執行 HCI IV DATA 轉檔 ===")
        self.trans_results.clear()

        search_pattern = os.path.join(target_dir, "*.[xlX][lsL]*")
        raw_files = [f for f in glob.glob(search_pattern) if not os.path.basename(f).startswith("~$")]

        if not raw_files:
            self.log("警告: 指定資料夾中找不到任何 Excel 檔案！")
            messagebox.showwarning("提示", "指定資料夾中找不到 Excel 檔案！")
            return

        self.log(f"找到 {len(raw_files)} 個原始 Excel 檔案進行處理。")

        try:
            results_folder = os.path.join(target_dir, "after_trans")
            os.makedirs(results_folder, exist_ok=True)

            con1_file = None

            for file_path in raw_files:
                source_file = os.path.basename(file_path)
                self.log(f"正在處理檔案: {source_file}")

                # 使用 openpyxl 直接讀取活頁簿（保留原始格式與大小寫分頁對應）
                wb = openpyxl.load_workbook(file_path, data_only=True)
                sheet_lower_map = {s.lower(): s for s in wb.sheetnames}

                required_sheets = ["Model team Fitting data", "IdVg_raw", "IdVd_raw", "Parameter"]
                real_sheets = {}
                for req in required_sheets:
                    req_lower = req.lower()
                    if req_lower in sheet_lower_map:
                        real_sheets[req] = sheet_lower_map[req_lower]
                    else:
                        raise ValueError(f"檔案 {source_file} 缺少必要的工作表: {req}")

                ws_fitting = wb[real_sheets["Model team Fitting data"]]
                ws_idvg = wb[real_sheets["IdVg_raw"]]
                ws_idvd = wb[real_sheets["IdVd_raw"]]
                ws_param = wb[real_sheets["Parameter"]]

                sp1 = source_file.split("_")
                if len(sp1) >= 2:
                    newname = f"{sp1[0]}_{sp1[1]}_IV_t0"
                    newname2 = f"{sp1[0]}_{sp1[1]}_IV_filnal_stress"
                else:
                    newname = f"{sp1[0]}_IV_t0"
                    newname2 = f"{sp1[0]}_IV_filnal_stress"

                lines_t0 = self.hci_IV_t0(ws_fitting, ws_idvg, ws_idvd)
                lines_final = self.hci_iv_final_stress(ws_fitting, ws_idvg, ws_idvd, ws_param)

                self.trans_results[newname] = lines_t0
                self.trans_results[newname2] = lines_final

                if "con1_" in source_file and not con1_file:
                    con1_file = file_path

            if con1_file:
                self.log(f"發現包含 con1_ 的檔案: {os.path.basename(con1_file)}，正在產生 isub_igate...")
                wb_con1 = openpyxl.load_workbook(con1_file, data_only=True)
                con1_lower_map = {s.lower(): s for s in wb_con1.sheetnames}
                
                if "ibig_vg" in con1_lower_map and "model team fitting data" in con1_lower_map:
                    ws_fitting_c1 = wb_con1[con1_lower_map["model team fitting data"]]
                    ws_ibig = wb_con1[con1_lower_map["ibig_vg"]]
                    lines_isub = self.isub_igate_iv(ws_fitting_c1, ws_ibig)
                    self.trans_results["isub_igate"] = lines_isub

            self.log("=== HCI IV DATA 轉檔完成！ ===")
            messagebox.showinfo("成功", "HCI IV 轉檔程序已全部完成！")

        except Exception as e:
            self.log(f"轉檔過程發生錯誤: {str(e)}")
            messagebox.showerror("錯誤", f"轉檔失敗:\n{str(e)}")

    def hci_IV_t0(self, ws_fitting, ws_idvg, ws_idvd):
        # openpyxl 的 row/col 從 1 開始計算 (對應 VBA row/col)
        w = self.get_cell_val(ws_fitting, 9, 2)
        l_val = self.get_cell_val(ws_fitting, 10, 2)
        t_val = self.get_cell_val(ws_fitting, 8, 2)
        vdlin = self.get_cell_val(ws_fitting, 14, 2)
        vdd = self.get_cell_val(ws_fitting, 13, 2)
        vdstr = self.get_cell_val(ws_fitting, 4, 2)
        vgstr = self.get_cell_val(ws_fitting, 3, 2)
        vbstr = self.get_cell_val(ws_fitting, 6, 2)
        time_str = "0"

        try:
            mos_val = float(ws_fitting.cell(row=3, column=2).value or 0)
            mos_type = "NMOS" if mos_val > 0 else "PMOS"
        except:
            mos_type = "NMOS"

        a = f"Instance{{W={w},L={l_val},T={t_val}}}"
        b = f"DataType{{{mos_type}}}"
        c = f"[vgs,Id(vbs=0.00000),vds={vdlin},ref_vs=0.00000,Vdstr={vdstr},Vgstr={vgstr},Vsstr=0,Vbstr={vbstr},Time={time_str}]"
        d = f"[vgs,Id(vbs=0.00000),vds={vdd},ref_vs=0.00000,Vdstr={vdstr},Vgstr={vgstr},Vsstr=0,Vbstr={vbstr},Time={time_str}]"
        e = f"[vds,Id(vgs={vdd}),vbs=0.00000,ref_vs=0.00000,Vdstr={vdstr},Vgstr={vgstr},Vsstr=0,Vbstr=0,Time={time_str}]"

        # 對應 VBA: G1 = Cells(15,2).End(xlDown).Row -> num = G1 - 15
        g1 = 15
        while g1 + 1 <= ws_idvg.max_row and ws_idvg.cell(row=g1 + 1, column=2).value is not None and str(ws_idvg.cell(row=g1 + 1, column=2).value).strip() != "":
            g1 += 1
        num = g1 - 15

        g2 = 18
        while g2 + 1 <= ws_idvd.max_row and ws_idvd.cell(row=g2 + 1, column=2).value is not None and str(ws_idvd.cell(row=g2 + 1, column=2).value).strip() != "":
            g2 += 1
        num2 = g2 - 18

        lines = [""] * (26 + 2 * num + num2)
        lines[0] = "//Copyright (C) 2010-2020 Primarius Technologies Co., Ltd"
        lines[1] = "// All rights reserved."
        lines[3] = "ObjInfo{}"
        lines[4] = "ProbeInfo{}"
        lines[5] = "Version{4.1}"
        lines[6] = "ModelType{DC}"
        lines[7] = b
        lines[8] = "Delimiter{,}"
        lines[9] = "WorkingMode{Forward}"
        lines[10] = a
        lines[11] = "Node{D|7|0.000000,G|8|0.000000,S|6|0.000000,B|5|0.000000}"
        lines[12] = "MeasurementSetup{PDA}"
        lines[13] = "Input{vgs,vbs,vds}"
        lines[14] = "Output{Id,Isub}"
        lines[15] = "{MOS_Id_Vgs_Vbs_DC}"
        lines[16] = c

        lines[18 + num] = "{MOS_Id_Vgs_Vbs_DC}"
        lines[19 + num] = d
        lines[23 + 2 * num] = "{MOS_Id_Vds_Vgs_DC}"
        lines[24 + 2 * num] = e

        for i in range(num + 1):
            row_idx = 15 + i
            v1 = self.get_cell_val(ws_idvg, row_idx, 1)
            v2 = self.get_cell_val(ws_idvg, row_idx, 2)
            v4 = self.get_cell_val(ws_idvg, row_idx, 4)
            lines[17 + i] = f"{v1},{v2}"
            lines[20 + num + i] = f"{v1},{v4}"

        for j in range(num2 + 1):
            row_idx = 18 + j
            v1 = self.get_cell_val(ws_idvd, row_idx, 1)
            v2 = self.get_cell_val(ws_idvd, row_idx, 2)
            lines[25 + 2 * num + j] = f"{v1},{v2}"

        return [line for line in lines if line is not None]

    def hci_iv_final_stress(self, ws_fitting, ws_idvg, ws_idvd, ws_param):
        w = self.get_cell_val(ws_fitting, 9, 2)
        l_val = self.get_cell_val(ws_fitting, 10, 2)
        t_val = self.get_cell_val(ws_fitting, 8, 2)
        vdlin = self.get_cell_val(ws_fitting, 14, 2)
        vdd = self.get_cell_val(ws_fitting, 13, 2)
        vdstr = self.get_cell_val(ws_fitting, 4, 2)
        vgstr = self.get_cell_val(ws_fitting, 3, 2)
        vbstr = self.get_cell_val(ws_fitting, 6, 2)

        # 尋找 Parameter sheet 中的 Tstress
        tstress_row = None
        for r in range(1, ws_param.max_row + 1):
            for c in range(1, ws_param.max_column + 1):
                val = ws_param.cell(row=r, column=c).value
                if val and str(val).strip().lower() == "tstress":
                    tstress_row = r
                    break
            if tstress_row is not None:
                break

        if tstress_row is None:
            raise ValueError("Parameter sheet 找不到 'Tstress'")

        g3 = tstress_row
        while g3 + 1 <= ws_param.max_row and ws_param.cell(row=g3 + 1, column=1).value is not None and str(ws_param.cell(row=g3 + 1, column=1).value).strip() != "":
            g3 += 1
        
        try:
            raw_time = float(ws_param.cell(row=g3, column=1).value or 0)
            time_str = str(raw_time / 60)
        except:
            time_str = self.get_cell_val(ws_param, g3, 1)

        try:
            mos_val = float(ws_fitting.cell(row=3, column=2).value or 0)
            mos_type = "NMOS" if mos_val > 0 else "PMOS"
        except:
            mos_type = "NMOS"

        a = f"Instance{{W={w},L={l_val},T={t_val}}}"
        b = f"DataType{{{mos_type}}}"
        c = f"[vgs,Id(vbs=0.00000),vds={vdlin},ref_vs=0.00000,Vdstr={vdstr},Vgstr={vgstr},Vsstr=0,Vbstr={vbstr},Time={time_str}]"
        d = f"[vgs,Id(vbs=0.00000),vds={vdd},ref_vs=0.00000,Vdstr={vdstr},Vgstr={vgstr},Vsstr=0,Vbstr={vbstr},Time={time_str}]"
        e = f"[vds,Id(vgs={vdd}),vbs=0.00000,ref_vs=0.00000,Vdstr={vdstr},Vgstr={vgstr},Vsstr=0,Vbstr=0,Time={time_str}]"

        g1 = 15
        while g1 + 1 <= ws_idvg.max_row and ws_idvg.cell(row=g1 + 1, column=2).value is not None and str(ws_idvg.cell(row=g1 + 1, column=2).value).strip() != "":
            g1 += 1
        num = g1 - 15

        g2 = 18
        while g2 + 1 <= ws_idvd.max_row and ws_idvd.cell(row=g2 + 1, column=2).value is not None and str(ws_idvd.cell(row=g2 + 1, column=2).value).strip() != "":
            g2 += 1
        num2 = g2 - 18

        g4 = 1
        while g4 + 1 <= ws_idvg.max_row and ws_idvg.cell(row=g4 + 1, column=1).value is not None and str(ws_idvg.cell(row=g4 + 1, column=1).value).strip() != "":
            g4 += 1
        num3 = max(0, g4 - num)

        g5 = 1
        while g5 + 1 <= ws_idvd.max_row and ws_idvd.cell(row=g5 + 1, column=1).value is not None and str(ws_idvd.cell(row=g5 + 1, column=1).value).strip() != "":
            g5 += 1
        num4 = max(0, g5 - num2)

        lines = [""] * (26 + 2 * num + num2)
        lines[0] = "//Copyright (C) 2010-2020 Primarius Technologies Co., Ltd"
        lines[1] = "// All rights reserved."
        lines[3] = "ObjInfo{}"
        lines[4] = "ProbeInfo{}"
        lines[5] = "Version{4.1}"
        lines[6] = "ModelType{DC}"
        lines[7] = b
        lines[8] = "Delimiter{,}"
        lines[9] = "WorkingMode{Forward}"
        lines[10] = a
        lines[11] = "Node{D|7|0.000000,G|8|0.000000,S|6|0.000000,B|5|0.000000}"
        lines[12] = "MeasurementSetup{PDA}"
        lines[13] = "Input{vgs,vbs,vds}"
        lines[14] = "Output{Id,Isub}"
        lines[15] = "{MOS_Id_Vgs_Vbs_DC}"
        lines[16] = c

        lines[18 + num] = "{MOS_Id_Vgs_Vbs_DC}"
        lines[19 + num] = d
        lines[23 + 2 * num] = "{MOS_Id_Vds_Vgs_DC}"
        lines[24 + 2 * num] = e

        for i in range(num + 1):
            target_idx = num3 + i
            v1 = self.get_cell_val(ws_idvg, target_idx, 1)
            v2 = self.get_cell_val(ws_idvg, target_idx, 2)
            v4 = self.get_cell_val(ws_idvg, target_idx, 4)
            lines[17 + i] = f"{v1},{v2}"
            lines[20 + num + i] = f"{v1},{v4}"

        for j in range(num2 + 1):
            target_idx = num4 + j
            v1 = self.get_cell_val(ws_idvd, target_idx, 1)
            v2 = self.get_cell_val(ws_idvd, target_idx, 2)
            lines[25 + 2 * num + j] = f"{v1},{v2}"

        return [line for line in lines if line is not None]

    def isub_igate_iv(self, ws_fitting, ws_ibig):
        w = self.get_cell_val(ws_fitting, 9, 2)
        l_val = self.get_cell_val(ws_fitting, 10, 2)
        t_val = self.get_cell_val(ws_fitting, 8, 2)
        vdd = self.get_cell_val(ws_ibig, 3, 2)
        
        try:
            mos_val = float(ws_fitting.cell(row=3, column=2).value or 0)
            mos_type = "NMOS" if mos_val > 0 else "PMOS"
        except:
            mos_type = "NMOS"

        a = f"Instance{{W={w},L={l_val},T={t_val}}}"
        b = f"DataType{{{mos_type}}}"

        g1 = 34
        while g1 + 1 <= ws_ibig.max_row and ws_ibig.cell(row=g1 + 1, column=1).value is not None and str(ws_ibig.cell(row=g1 + 1, column=1).value).strip() != "":
            g1 += 1
        vg_num = g1 - 34

        lines = [""] * (22 + vg_num * 2)
        lines[0] = "//Copyright (C) 2010-2020 Primarius Technologies Co., Ltd"
        lines[1] = "// All rights reserved."
        lines[3] = "ObjInfo{}"
        lines[4] = "ProbeInfo{}"
        lines[5] = "Version{4.1}"
        lines[6] = "ModelType{DC}"
        lines[7] = b
        lines[8] = "Delimiter{,}"
        lines[9] = "WorkingMode{Forward}"
        lines[10] = a
        lines[11] = "Node{D|7|0.000000,G|8|0.000000,S|6|0.000000,B|5|0.000000}"
        lines[12] = "MeasurementSetup{PDA}"
        lines[13] = "Input{vgs,vbs,vds}"
        lines[14] = "Output{Id,Isub}"
        lines[15] = "{MOS_Isub_Vgs_Vds_DC}"
        lines[16] = f"[vgs,Isub(vds={vdd}),vbs=0.00000,ref_vs=0.00000]"
        lines[18 + vg_num] = "{Igate_Vgs_Vds_DC}"
        lines[19 + vg_num] = f"[vgs,Ig(vds={vdd}),vbs=0.00000,ref_vs=0.00000]"

        for i in range(vg_num + 1):
            row_idx = 34 + i
            v1 = self.get_cell_val(ws_ibig, row_idx, 1)
            v5 = self.get_cell_val(ws_ibig, row_idx, 5)
            v3 = self.get_cell_val(ws_ibig, row_idx, 3)
            lines[17 + i] = f"{v1},{v5}"
            lines[20 + vg_num + i] = f"{v1},{v3}"

        return [line for line in lines if line is not None]

    def run_save_dat(self):
        """執行 SAVE IV DATA 按鈕邏輯"""
        target_dir = self.path_entry.get().strip()
        if not target_dir or not os.path.isdir(target_dir):
            messagebox.showerror("錯誤", "請先選擇有效的資料夾路徑！")
            return

        if not self.trans_results:
            messagebox.showwarning("提示", "目前沒有可儲存的轉檔資料，請先執行 'HCI IV DATA' 轉檔！")
            return

        self.log("=== 開始儲存 .dat 檔案 ===")
        results_folder = os.path.join(target_dir, "after_trans")
        os.makedirs(results_folder, exist_ok=True)

        saved_count = 0
        for sheet_name, lines in self.trans_results.items():
            file_path = os.path.join(results_folder, f"{sheet_name}.dat")
            try:
                with open(file_path, "w", encoding="utf-8") as f:
                    for line in lines:
                        f.write(str(line) + "\n")
                self.log(f"已儲存: {sheet_name}.dat")
                saved_count += 1
            except Exception as e:
                self.log(f"儲存 {sheet_name}.dat 失敗: {str(e)}")

        self.log(f"=== 儲存完成！共輸出 {saved_count} 個檔案至 after_trans 資料夾 ===")
        messagebox.showinfo("成功", f"成功儲存 {saved_count} 個 .dat 檔案到 after_trans 目錄！")

if __name__ == "__main__":
    root = tk.Tk()
    app = IVTransApp(root)
    root.mainloop()
